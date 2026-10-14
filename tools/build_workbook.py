#!/usr/bin/env python3
"""產生試用 Excel 活頁簿。
用法：node tools/export-workbook.js [--clean] > data.json
      python3 tools/build_workbook.py data.json 輸出.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="DCEBE8")
TITLE = Font(name=FONT, size=14, bold=True)
BOLD = Font(name=FONT, size=10, bold=True)
BODY = Font(name=FONT, size=10)
INPUT = Font(name=FONT, size=10, color="0000FF")      # 輸入值
NOTE = Font(name=FONT, size=10, color="5A6873")

data = json.load(open(sys.argv[1], encoding="utf-8"))
out = sys.argv[2]
S = data["sheets"]
wb = Workbook()


def header(ws, cols, widths=None):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for i, col in enumerate(cols, 1):
        w = (widths or {}).get(col) or max(8, min(28, len(str(col)) * 2 + 2))
        ws.column_dimensions[get_column_letter(i)].width = w


def text_sheet(name, sheet, widths=None):
    """所有欄位存成文字，保留前置 0 與空白 (固定長度欄位)。"""
    ws = wb.create_sheet(name)
    header(ws, sheet["header"], widths)
    for row in sheet["rows"]:
        ws.append([str(v) for v in row])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.number_format, c.font = "@", BODY
    ws.auto_filter.ref = ws.dimensions
    return ws


# ---------- 說明 ----------
ws = wb.active
ws.title = "說明"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 80
ws["A1"] = "申報資料檢核助手｜試用資料" + ("（含錯誤）" if data["inject"] else "（正確）")
ws["A1"].font = TITLE
intro = [
    f"資料組：seed={data['dataset']['seed']}, policies={data['dataset']['policies']}（tests/run.html 用來對照錯誤解答）",
    "用途：個人壽險年度申報的模擬資料，供試用「申報資料檢核助手」。資料皆為假設，不含真實個資。",
    "使用方式：開啟檢核工具，把這個 Excel 檔拖進上方的大框，工具會依工作表名稱自動分好。",
    ("本檔刻意放入各種錯誤，答案在「錯誤解答」工作表；試用時可先不看答案，比對工具找到的結果。"
     if data["inject"] else "本檔是正確資料，檢查結果應該是 0 個問題，可用來和含錯誤版本對照。"),
    "欄位：保單、理賠、去年保單資料的欄位依《人壽保險業個人壽險填報手冊》(2025 填報年度)；數字以文字儲存，保留開頭的 0。",
]
for i, t in enumerate(intro, 3):
    ws.cell(row=i, column=1, value=t).font = BODY
row = 3 + len(intro) + 1
ws.cell(row=row, column=1, value="工作表").font = BOLD
ws.cell(row=row, column=2, value="筆數").font = BOLD
ws.cell(row=row, column=3, value="內容").font = BOLD
for c in range(1, 4):
    ws.cell(row=row, column=c).fill = HEAD_FILL
desc = {
    "保單資料": "每張保單的申報資料（61 欄）",
    "理賠資料": "每筆理賠給付的申報資料（31 欄）",
    "客戶資料": "要保人（角色 1）與被保險人（角色 2）",
    "商品資料": "商品代號、險種、繳費年期、起售日、停售日、替代代號",
    "去年保單資料": "去年申報的保單資料，用來確認去年底和今年初接得上",
    "系統總數": "新契約、理賠、保單系統各自的筆數與金額總數，用來核對申報資料",
    "資產負債表": "各科目餘額，上層科目用公式加總",
    "錯誤解答": "刻意放入的錯誤和工具應找到的結果（含 4 筆「請人工判斷」的不太合理案例）",
}
order = ["保單資料", "理賠資料", "客戶資料", "商品資料", "去年保單資料", "系統總數", "資產負債表"]
if data["inject"]:
    order.append("錯誤解答")
for name in order:
    row += 1
    ws.cell(row=row, column=1, value=name).font = BODY
    ws.cell(row=row, column=2, value=f"=COUNTA('{name}'!A:A)-1").font = BODY
    ws.cell(row=row, column=3, value=desc[name]).font = BODY

row += 2
ws.cell(row=row, column=1, value="資產負債表平衡檢查").font = BOLD
bs = S["資產負債表"]["rows"]
r1 = next(i for i, r in enumerate(bs) if r[0] == "1") + 2
r2 = next(i for i, r in enumerate(bs) if r[0] == "2") + 2
r3 = next(i for i, r in enumerate(bs) if r[0] == "3") + 2
for label, col in (("期初：資產 −（負債 + 權益）", "C"), ("期末：資產 −（負債 + 權益）", "F")):
    row += 1
    ws.cell(row=row, column=1, value=label).font = BODY
    c = ws.cell(row=row, column=2, value=f"='資產負債表'!{col}{r1}-'資產負債表'!{col}{r2}-'資產負債表'!{col}{r3}")
    c.font, c.number_format = BODY, '#,##0.00;(#,##0.00);"-"'
    ws.cell(row=row, column=3, value="不為 0 表示資產負債表不平").font = NOTE
for t in data["skipped"]:
    row += 2
    ws.cell(row=row, column=1, value="注意：" + t).font = NOTE
row += 2
ws.cell(row=row, column=1, value="圖例：資產負債表中藍字為輸入值，黑字為公式。").font = NOTE

# ---------- 資料工作表 ----------
text_sheet("保單資料", S["保單資料"])
text_sheet("理賠資料", S["理賠資料"])
text_sheet("客戶資料", S["客戶資料"], {"地址": 34, "姓名": 10, "證件編號": 14, "保單編號": 18})
text_sheet("商品資料", S["商品資料"], {"商品代號": 20, "商品名稱": 26, "替代代號": 20})
text_sheet("去年保單資料", S["去年保單資料"])

ws = wb.create_sheet("系統總數")
header(ws, S["系統總數"]["header"], {"項目名稱": 20, "來源系統": 14, "金額": 18})
for r in S["系統總數"]["rows"]:
    ws.append([r[0], r[1], r[2], r[3], int(r[4]), float(r[5])])
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = BODY
    r[3].number_format = "@"
    r[4].number_format = "#,##0"
    r[5].number_format = "#,##0.00"

ws = wb.create_sheet("資產負債表")
st = S["資產負債表"]
header(ws, st["header"], {"科目名稱": 40, "期初餘額": 18, "本期增加": 18, "本期減少": 18, "期末餘額": 18})
for i, r in enumerate(st["rows"]):
    ws.append([r[0], r[1]] + r[2:])
    excel_row = i + 2
    kids = st["children"][i]
    ws.cell(row=excel_row, column=1).number_format = "@"
    for col in range(3, 7):
        c = ws.cell(row=excel_row, column=col)
        c.number_format = "#,##0.00"
        if kids:   # 上層科目：下層合計
            L = get_column_letter(col)
            c.value = "=SUM(" + ",".join(f"{L}{k + 2}" for k in kids) + ")"
            c.font = BOLD if len(r[0]) == 1 else BODY
        else:
            c.font = INPUT
    for col in (1, 2):
        ws.cell(row=excel_row, column=col).font = BOLD if len(r[0]) <= 2 else BODY
    ws.cell(row=excel_row, column=2).alignment = Alignment(indent=0 if len(r[0]) == 1 else 1 if len(r[0]) == 2 else 2)

if data["inject"]:
    ws = wb.create_sheet("錯誤解答")
    header(ws, ["類別", "規則編號", "工作表", "保單號碼／科目／期間", "Excel 列號", "說明"],
           {"類別": 22, "規則編號": 16, "工作表": 18, "保單號碼／科目／期間": 24, "Excel 列號": 10, "說明": 60})
    for k in data["key"]:
        ws.append(k)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font = BODY
    ws.auto_filter.ref = ws.dimensions

wb.save(out)
print(out)
