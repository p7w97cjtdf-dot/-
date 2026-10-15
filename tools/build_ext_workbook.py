#!/usr/bin/env python3
"""產生癌症險 / 傷害險的試用 Excel 活頁簿。
用法：node tools/build-ext-samples.js cancer
      python3 tools/build_ext_workbook.py /tmp/ext-cancer-dirty.json 輸出.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="DCEBE8")
TITLE = Font(name=FONT, size=14, bold=True)
BOLD = Font(name=FONT, size=10, bold=True)
BODY = Font(name=FONT, size=10)

data = json.load(open(sys.argv[1], encoding="utf-8"))
S = data["sheets"]
T = data["title"]
MANUAL = {"cancer": "《人壽保險業個人癌症保險填報手冊》", "accident": "《人壽保險業新式傷害保險填報手冊》"}[data["product"]]
wb = Workbook()


def header(ws, cols, widths=None):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col) or max(8, min(28, len(str(col)) * 2 + 2))


ws = wb.active
ws.title = "說明"
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 80
ws["A1"] = f"申報資料檢核助手｜{T}試用資料" + ("（含錯誤）" if data["inject"] else "（正確）")
ws["A1"].font = TITLE
intro = [
    f"用途：個人{T}年度申報（填報年度 {data['year']}）的模擬資料，資料皆為假設，不含真實個資。",
    "使用方式：開啟檢核工具，把這個 Excel 檔拖進上方的大框，工具會依工作表名稱自動分好。",
    ("本檔刻意放入各種錯誤，答案在「錯誤解答」工作表；試用時可先不看答案，比對工具找到的結果。"
     if data["inject"] else "本檔是正確資料，檢查結果應該是 0 個問題，可用來和含錯誤版本對照。"),
    f"欄位依{MANUAL}；數字以文字儲存，保留開頭的 0 與空白。",
]
for i, t in enumerate(intro, 3):
    ws.cell(row=i, column=1, value=t).font = BODY
row = 3 + len(intro) + 1
for c, v in enumerate(["工作表", "筆數", "內容"], 1):
    cell = ws.cell(row=row, column=c, value=v)
    cell.font, cell.fill = BOLD, HEAD_FILL
names = list(S.keys())
desc = {names[0]: f"每張保單的承保資料（{data['layout']['L']} 欄）", names[1]: f"每筆給付的資料（{data['layout']['D']} 欄）", "錯誤解答": "刻意放入的錯誤和工具應找到的檢核編號"}
for name in names + (["錯誤解答"] if data["inject"] else []):
    row += 1
    ws.cell(row=row, column=1, value=name).font = BODY
    ws.cell(row=row, column=2, value=f"=COUNTA('{name}'!A:A)-1").font = BODY
    ws.cell(row=row, column=3, value=desc[name]).font = BODY

for name, sh in S.items():
    w = wb.create_sheet(name)
    header(w, sh["header"])
    for r in sh["rows"]:
        w.append([str(v) for v in r])
    for r in w.iter_rows(min_row=2):
        for c in r:
            c.number_format, c.font = "@", BODY
    w.auto_filter.ref = w.dimensions

if data["inject"]:
    w = wb.create_sheet("錯誤解答")
    header(w, ["應找到的檢核編號", "工作表", "保單號碼", "Excel 列", "刻意放入的錯誤"], {"刻意放入的錯誤": 60, "應找到的檢核編號": 26})
    for k in data["key"]:
        w.append(k)
    for r in w.iter_rows(min_row=2):
        for c in r:
            c.font = BODY
wb.save(sys.argv[2])
